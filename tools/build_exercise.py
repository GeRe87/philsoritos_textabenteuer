from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
import matplotlib.pyplot as plt
import math
import os

OUT = "build"
os.makedirs(OUT, exist_ok=True)

# Diagram: global IUCN estimate for primary microplastic sources
labels = [
    "Synthetische Textilien", "Reifenabrieb", "Stadtstaub",
    "Straßenmarkierungen", "Schiffsbeschichtungen",
    "Körperpflege", "Kunststoffpellets"
]
values = [35, 28, 24, 7, 3.7, 2, 0.3]
fig, ax = plt.subplots(figsize=(8, 4.6))
ax.barh(labels[::-1], values[::-1])
ax.set_xlabel("Anteil an primären Mikroplastikeinträgen ins Meer [%]")
ax.set_title("Globale Schätzung nach IUCN (2017)")
ax.set_xlim(0, 40)
for i, value in enumerate(values[::-1]):
    ax.text(value + 0.5, i, f"{value:g} %", va="center", fontsize=9)
fig.tight_layout()
chart_path = os.path.join(OUT, "diagramm_quellen.png")
fig.savefig(chart_path, dpi=180, bbox_inches="tight")
plt.close(fig)

# Original schematic illustration
W, H = 1400, 750
img = Image.new("RGB", (W, H), (240, 246, 249))
d = ImageDraw.Draw(img)
try:
    font_big = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 44)
    font_small = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 23)
except OSError:
    font_big = font_small = None

for y in range(H):
    d.line((0, y, W, y), fill=(230, 242, 249) if y < 430 else (222, 229, 214))
d.rectangle((0, 560, W, H), fill=(185, 220, 235))
d.text((45, 35), "Vom Alltag ins Gewässer", fill=(30, 60, 75), font=font_big)

# House + washing machine
hx, hy = 90, 230
d.rectangle((hx, hy, hx + 260, hy + 250), fill=(250, 248, 240), outline=(80, 80, 80), width=4)
d.polygon([(hx - 20, hy), (hx + 130, hy - 100), (hx + 280, hy)], fill=(190, 110, 85), outline=(80, 80, 80))
d.rectangle((hx + 65, hy + 85, hx + 195, hy + 205), fill=(225, 228, 230), outline=(80, 80, 80), width=3)
d.ellipse((hx + 90, hy + 110, hx + 170, hy + 190), fill=(170, 205, 220), outline=(70, 90, 100), width=3)
d.text((hx + 42, hy + 215), "synthetische Textilien", fill=(40, 40, 40), font=font_small)

# Road + car
d.rectangle((410, 455, 900, 550), fill=(120, 125, 130))
d.line((430, 502, 880, 502), fill=(245, 240, 190), width=8)
carx, cary = 520, 365
d.rounded_rectangle((carx, cary, carx + 270, cary + 105), radius=22, fill=(205, 70, 70), outline=(80, 60, 60), width=4)
d.polygon([(carx + 55, cary), (carx + 105, cary - 70), (carx + 200, cary - 70), (carx + 240, cary)], fill=(205, 70, 70), outline=(80, 60, 60))
for wx in (carx + 65, carx + 220):
    d.ellipse((wx - 28, cary + 78, wx + 28, cary + 134), fill=(35, 35, 35), outline=(10, 10, 10))
d.text((500, 330), "Reifenabrieb", fill=(40, 40, 40), font=font_small)

# City
for x, h in [(960, 170), (1040, 240), (1140, 195), (1230, 280)]:
    d.rectangle((x, 560 - h, x + 70, 560), fill=(165, 175, 182), outline=(90, 100, 105))
    for yy in range(560 - h + 20, 550, 40):
        for xx in range(x + 14, x + 60, 24):
            d.rectangle((xx, yy, xx + 10, yy + 15), fill=(245, 235, 165))
d.text((965, 250), "Stadtstaub & Abrieb", fill=(40, 40, 40), font=font_small)
d.rectangle((870, 500, 940, 590), fill=(110, 110, 115), outline=(70, 70, 70))
d.line((900, 590, 900, 650), fill=(80, 110, 130), width=10)

for x1, y1, x2, y2 in [(300, 480, 650, 600), (740, 500, 820, 610), (1110, 540, 970, 620)]:
    d.line((x1, y1, x2, y2), fill=(55, 120, 155), width=8)
    angle = math.atan2(y2 - y1, x2 - x1)
    for a in (angle + 2.6, angle - 2.6):
        d.line((x2, y2, x2 + 24 * math.cos(a), y2 + 24 * math.sin(a)), fill=(55, 120, 155), width=8)

for x, y, r, c in [
    (250, 640, 8, (210, 70, 80)), (390, 690, 6, (245, 170, 60)),
    (590, 620, 7, (105, 85, 180)), (760, 700, 9, (60, 140, 100)),
    (930, 655, 5, (210, 70, 80)), (1120, 680, 8, (245, 170, 60)),
    (1280, 625, 6, (105, 85, 180))
]:
    d.ellipse((x - r, y - r, x + r, y + r), fill=c, outline=(70, 70, 70))
d.text((40, 690), "Schematische Illustration: Partikel gelangen über verschiedene Wege in Böden, Kanalisation und Gewässer.", fill=(35, 55, 65), font=font_small)
illustration_path = os.path.join(OUT, "bild_wege.png")
img.save(illustration_path)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(2.1)
sec.left_margin = Cm(2.1)
sec.right_margin = Cm(2.1)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Seite ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)

add_page_number(sec.footer.paragraphs[0])
fonts = ["Arial", "Times New Roman", "Calibri", "Verdana"]

def add_heading(text, level=1, variant=0):
    p = doc.add_paragraph()
    if variant == 0:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.bold = True; r.font.name = "Arial"; r.font.size = Pt(19 if level == 1 else 15)
    elif variant == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text); r.bold = True; r.underline = True; r.font.name = "Times New Roman"; r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(40, 85, 140)
    elif variant == 2:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text.upper()); r.bold = True; r.font.name = "Verdana"; r.font.size = Pt(13)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(text); r.italic = True; r.font.name = "Calibri"; r.font.size = Pt(17)
    p.paragraph_format.space_before = Pt([0, 8, 3, 14][variant % 4])
    p.paragraph_format.space_after = Pt([8, 2, 11, 4][variant % 4])

def add_body(text, variant=0):
    p = doc.add_paragraph()
    p.alignment = [
        WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.LEFT,
        WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER
    ][variant % 4]
    p.paragraph_format.line_spacing = [1.0, 1.35, 1.15, 1.0][variant % 4]
    p.paragraph_format.space_after = Pt([2, 8, 0, 5][variant % 4])
    if variant == 2:
        p.paragraph_format.first_line_indent = Cm(0.8)
    r = p.add_run(text)
    r.font.name = fonts[variant % len(fonts)]
    r.font.size = Pt([10.5, 11.5, 10, 11][variant % 4])

# Page 1
add_heading("Mikroplastik – unsichtbare Partikel im Alltag und in der Umwelt", 1, 0)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run("Formatierungsübung – Rohfassung")
r.bold = True; r.font.name = "Courier New"; r.font.size = Pt(11)

add_body("""Kunststoffe sind aus dem modernen Alltag kaum wegzudenken. Sie werden für Verpackungen, Kleidung, Fahrzeuge, Medizinprodukte, Elektronik und zahlreiche weitere Anwendungen eingesetzt. Ihre Vorteile liegen unter anderem in geringem Gewicht, Formbarkeit, Beständigkeit und vergleichsweise niedrigen Herstellungskosten. Genau diese Beständigkeit wird jedoch zum Problem, wenn Kunststoffprodukte in die Umwelt gelangen. Größere Gegenstände können durch Sonnenlicht, Temperaturwechsel, mechanische Belastung und Wellenbewegung in immer kleinere Fragmente zerfallen. Daneben entstehen sehr kleine Kunststoffpartikel bereits während der Nutzung bestimmter Produkte, etwa durch Reifenabrieb oder beim Waschen synthetischer Textilien. Solche Partikel werden im Allgemeinen als Mikroplastik bezeichnet, wenn sie kleiner als fünf Millimeter sind. Die genaue Abgrenzung zu noch kleineren Nanoplastikpartikeln ist je nach Fachgebiet nicht völlig einheitlich, weshalb bei Vergleichen zwischen Studien immer auch die verwendete Definition betrachtet werden muss.""", 0)

add_body("""Mikroplastik ist nicht auf Meere beschränkt. Untersuchungen weisen Kunststoffpartikel in Flüssen, Seen, Sedimenten, Böden, Abwässern und in der Luft nach. Die Partikel unterscheiden sich in Größe, Form, Polymerart und chemischen Zusatzstoffen. Fasern aus Polyester oder Polyamid sehen unter dem Mikroskop anders aus als unregelmäßige Fragmente einer gealterten Verpackung oder kugelförmige Kunststoffpellets. Diese Vielfalt erschwert die Analytik: Für Probenahme, Aufbereitung, Identifizierung und Mengenbestimmung existieren unterschiedliche Verfahren, die nicht immer direkt miteinander vergleichbar sind. Deshalb ist es wichtig, Messergebnisse nicht nur als einzelne Zahl zu betrachten, sondern auch zu fragen, welche Partikelgrößen erfasst wurden und mit welcher Methode die Identifikation erfolgte.""", 1)

add_heading("1  Was ist primäres und sekundäres Mikroplastik?", 2, 1)
add_body("""Häufig wird zwischen primärem und sekundärem Mikroplastik unterschieden. Primäres Mikroplastik gelangt bereits als kleines Partikel in die Umwelt oder wird während der Nutzung eines Produkts in dieser Größe erzeugt. Dazu gehören beispielsweise Fasern aus synthetischen Textilien, Partikel aus Reifen- und Straßenabrieb sowie bestimmte industriell hergestellte Pellets. Sekundäres Mikroplastik entsteht dagegen aus größeren Kunststoffteilen. Eine Plastiktüte, eine Getränkeflasche oder ein Stück Verpackungsfolie kann durch Verwitterung verspröden und nach und nach in kleinere Fragmente zerfallen. Diese Unterscheidung ist nützlich, weil sich mögliche Minderungsmaßnahmen unterscheiden: Bei sekundärem Mikroplastik spielt die Vermeidung von Kunststoffabfällen und Littering eine große Rolle, während bei primären Quellen häufig Produktdesign, Materialwahl, Abrieb und technische Rückhaltesysteme betrachtet werden müssen.""", 2)

add_body("""Dass Mikroplastik heute fast überall gesucht und häufig gefunden wird, bedeutet nicht automatisch, dass jede nachgewiesene Menge unmittelbar ein Gesundheitsrisiko darstellt. Für eine Risikobewertung müssen Exposition, Partikeleigenschaften und biologische Wirkung gemeinsam betrachtet werden. Die Weltgesundheitsorganisation betonte bereits in ihrem Bericht zu Mikroplastik im Trinkwasser, dass Datenlücken und methodische Unterschiede eine vorsichtige Interpretation erfordern. Gleichzeitig ist die weite Verbreitung ein starkes Argument dafür, Einträge zu reduzieren und Messmethoden weiter zu standardisieren.""", 3)
doc.add_page_break()

# Page 2
add_heading("Woher kommen die Partikel?", 1, 2)
add_body("""Ein Teil der Mikroplastikbelastung entsteht durch alltägliche Nutzung. Beim Fahren wird Material von Reifen und Straßenoberflächen abgerieben. Beim Waschen synthetischer Kleidung können feine Fasern freigesetzt werden. Farben und Beschichtungen können altern und Partikel verlieren. In Städten kommen weitere Abrieb- und Staubquellen hinzu. Gleichzeitig zerfällt achtlos entsorgter oder schlecht bewirtschafteter Kunststoffabfall in der Umwelt über lange Zeiträume. Die Einträge gelangen über Regenwasser, Kanalisation, Flüsse, Kläranlagen, Windtransport oder direkte Freisetzung in verschiedene Umweltkompartimente.""", 1)

add_heading("Tabelle 1: Beispiele wichtiger Entstehungs- und Eintragspfade", 2, 3)
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
table.autofit = False
widths = [Cm(3.1), Cm(4.0), Cm(4.0), Cm(4.2)]
headers = ["Quelle", "Wie entstehen Partikel?", "Möglicher Eintragspfad", "Mögliche Ansatzpunkte"]
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.width = widths[i]
    cell.text = h
    for rr in cell.paragraphs[0].runs:
        rr.bold = True; rr.font.name = "Arial"; rr.font.size = Pt(9)

rows = [
    ("Synthetische Textilien", "Faserabrieb beim Tragen und Waschen", "Abwasser, Hausstaub, Luft", "Materialdesign, Waschverhalten, Filtertechnik"),
    ("Fahrzeugreifen", "Abrieb durch Reibung auf der Straße", "Straßenabfluss, Boden, Luft", "Reifenentwicklung, Verkehr, Rückhalt von Straßenabfluss"),
    ("Größere Kunststoffabfälle", "Verwitterung und mechanischer Zerfall", "Boden, Flüsse, Küsten, Meer", "Abfallvermeidung, Sammlung, Recycling"),
    ("Farben und Beschichtungen", "Alterung, Abplatzen, Schleifen", "Boden, Wasser, Staub", "Beständigere Systeme, Auffangen bei Sanierung"),
    ("Kunststoffpellets", "Verlust bei Produktion und Transport", "Industriegelände, Kanalisation, Flüsse", "Verlustprävention und sauberes Handling"),
]
for ri, row in enumerate(rows):
    cells = table.add_row().cells
    for ci, value in enumerate(row):
        cells[ci].width = widths[ci]
        cells[ci].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        pp = cells[ci].paragraphs[0]
        pp.alignment = WD_ALIGN_PARAGRAPH.LEFT if (ri + ci) % 2 else WD_ALIGN_PARAGRAPH.CENTER
        rr = pp.add_run(value)
        rr.font.name = ["Calibri", "Times New Roman", "Arial", "Verdana"][(ri + ci) % 4]
        rr.font.size = Pt(8.5 if ci != 0 else 9.5)

add_body("""Die Tabelle zeigt, dass Mikroplastik kein einzelnes Produktproblem ist. Manche Quellen lassen sich direkt am Produkt beeinflussen, andere hängen mit Infrastruktur, Verkehrsmenge oder Abfallwirtschaft zusammen. Auch Kläranlagen spielen eine doppelte Rolle. Sie können einen großen Anteil der Partikel aus dem Abwasser entfernen, doch diese Partikel verschwinden dadurch nicht. Sie werden in Rückständen wie Klärschlamm angereichert. Je nachdem, wie solche Rückstände weiterbehandelt oder genutzt werden, können Partikel in andere Umweltbereiche verlagert werden. Daher sollte eine Bewertung möglichst den gesamten Stoffstrom betrachten und nicht nur die Konzentration im gereinigten Wasser.""", 0)

add_body("""Internationale Abschätzungen zeigen außerdem, dass die relative Bedeutung einzelner Quellen von Region und betrachteter Umwelt abhängt. In einer häufig zitierten globalen IUCN-Abschätzung für primäres Mikroplastik im Meer entfielen große Anteile auf synthetische Textilien, Reifenabrieb und Stadtstaub. Solche Prozentwerte sind jedoch keine universellen Naturkonstanten. Sie beruhen auf Modellannahmen und verfügbaren Daten und sollten deshalb als Größenordnung verstanden werden. Neuere europäische Arbeiten weisen ebenfalls darauf hin, dass Reifen, Farben, Textilien und Kunststoffpellets wichtige unbeabsichtigte Emissionsquellen sein können.""", 2)
doc.add_page_break()

# Page 3
add_heading("2. Vom Entstehungsort bis in Flüsse und Meere", 1, 3)
add_body("""Nach der Freisetzung können Mikroplastikpartikel verschiedene Wege nehmen. Größere und dichtere Partikel setzen sich eher in Böden oder Sedimenten ab, während leichte Fasern und kleine Fragmente über längere Strecken transportiert werden können. Niederschlag kann Partikel von Straßen und versiegelten Flächen in die Kanalisation spülen. In Mischwassersystemen können Starkregenereignisse zusätzliche Einträge verursachen. Flüsse verbinden schließlich viele landbasierte Quellen mit Seen, Küsten und Meeren. Gleichzeitig kann Wind sehr kleine Partikel transportieren und damit Regionen erreichen, die weit von offensichtlichen Kunststoffquellen entfernt liegen.""", 0)

add_heading("Diagramm: geschätzte Anteile primärer Mikroplastikquellen", 2, 1)
doc.add_picture(chart_path, width=Cm(13.0))
cap = doc.add_paragraph("Abbildung ?  Quelle: IUCN 2017 – Zahlen global modelliert, nicht als Messwert für einen einzelnen Ort zu verstehen.")
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
for rr in cap.runs:
    rr.font.name = "Courier New"; rr.font.size = Pt(8); rr.italic = True

add_body("""Die dargestellte Verteilung macht sichtbar, warum einfache Verbraucherbotschaften wie „verzichte auf Mikroperlen“ nur einen Teil des Problems erfassen. Absichtlich zugesetzte Mikroplastikpartikel in einzelnen Produkten können regulatorisch relativ gezielt adressiert werden. Abriebquellen entstehen dagegen während der Nutzung von Produkten und sind technisch komplexer. Die Europäische Union hat für absichtlich zugesetzte Mikroplastikpartikel bereits eine REACH-Beschränkung eingeführt. Parallel werden unbeabsichtigte Freisetzungen, beispielsweise aus Reifen, Textilien, Farben und Pellets, über weitere Maßnahmen adressiert. Für Schülerinnen und Schüler ist wichtig: Regulierung kann an unterschiedlichen Punkten ansetzen – bei der Herstellung, beim Produktdesign, bei der Nutzung, bei Filtern oder Rückhaltesystemen und bei der Abfallbehandlung.""", 1)

add_body("""Für wissenschaftliche Aussagen müssen außerdem Einheiten und Bezugsgrößen sauber angegeben werden. Manche Studien berichten Partikel pro Liter Wasser, andere Partikel pro Kilogramm Sediment oder Massekonzentrationen. Ein Vergleich ist schwierig, wenn unterschiedliche Größenklassen, Probenvolumina und analytische Nachweisgrenzen verwendet werden. Dasselbe gilt für Diagramme: Ein Balkendiagramm kann eine klare Botschaft vermitteln, aber nur dann korrekt interpretiert werden, wenn Quelle, Einheit, Bezugsraum und Unsicherheit genannt werden. Gerade deshalb gehören aussagekräftige Abbildungsunterschriften und vollständige Quellenangaben zu einer guten wissenschaftlichen Formatierung.""", 2)
doc.add_page_break()

# Page 4
add_heading("Auswirkungen auf Organismen und Ökosysteme", 1, 1)
add_body("""Mikroplastik kann von Organismen aufgenommen werden, wenn Partikel mit Nahrung verwechselt werden oder zusammen mit Wasser in den Verdauungstrakt gelangen. Beobachtet wurde die Aufnahme bei zahlreichen aquatischen Organismen. Welche Folgen daraus entstehen, hängt von Partikelgröße, Form, Konzentration, Polymer, Zusatzstoffen und der jeweiligen Art ab. In Laborstudien werden unter bestimmten Bedingungen unter anderem Veränderungen von Nahrungsaufnahme, Wachstum oder physiologischen Stressreaktionen untersucht. Die Übertragung solcher Ergebnisse auf reale Umweltbedingungen ist jedoch anspruchsvoll, weil experimentelle Konzentrationen und Partikeltypen nicht immer denen in der Umwelt entsprechen.""", 3)

add_body("""Neben möglichen direkten Partikeleffekten wird diskutiert, dass Kunststoffoberflächen andere Stoffe anlagern oder Mikroorganismen tragen können. Kunststoffe enthalten zudem Zusatzstoffe, die je nach Material und Alterungsverlauf freigesetzt werden können. Daraus folgt nicht, dass jedes Mikroplastikpartikel automatisch eine hochtoxische „Giftkugel“ ist. Für eine belastbare Bewertung müssen Konzentrationen, chemische Zusammensetzung, Bioverfügbarkeit und reale Expositionsdauer betrachtet werden. Genau an dieser Stelle zeigt sich die Bedeutung sauberer naturwissenschaftlicher Sprache: Zwischen einem nachgewiesenen Vorkommen, einem experimentell beobachteten Effekt und einem nachgewiesenen Risiko für eine Population besteht ein Unterschied.""", 0)

add_heading("Bild 1 – mögliche Wege aus dem Alltag", 2, 2)
doc.add_picture(illustration_path, width=Cm(13.2))
cap = doc.add_paragraph("grafik eigene darstellung / muss noch richtig beschriftet werden")
cap.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for rr in cap.runs:
    rr.font.name = "Comic Sans MS"; rr.font.size = Pt(9); rr.bold = True

add_body("""Auch die räumliche Verteilung ist nicht gleichmäßig. Strömung, Wind, Sedimentation und lokale Einträge führen dazu, dass bestimmte Bereiche stärker belastet sein können als andere. An Küsten können sich Partikel in Sedimenten anreichern, während leichte Partikel in der Wassersäule oder an der Oberfläche transportiert werden. Auf dem Land können Mikroplastikpartikel über Reifenabrieb, Kompost, Klärschlamm oder Folienreste in Böden gelangen. Damit ist Mikroplastik ein Beispiel für ein Umweltproblem, das Wasser, Boden und Luft verbindet und deshalb nicht sinnvoll in nur einem Umweltkompartiment betrachtet werden kann.""", 1)

# Page 5 flows naturally
add_heading("3 Gesundheit, Unsicherheit und sinnvolle Maßnahmen", 1, 0)
add_body("""Menschen können Mikroplastik über Nahrung und Trinkwasser aufnehmen und Partikel aus der Luft einatmen. In den vergangenen Jahren wurden Mikroplastikpartikel auch in verschiedenen menschlichen Proben beschrieben. Solche Nachweise zeigen Exposition, beantworten aber nicht automatisch die Frage, welche gesundheitliche Bedeutung die gefundenen Mengen haben. Reviews diskutieren mögliche Mechanismen wie Entzündungsreaktionen oder oxidativen Stress, weisen zugleich aber auf große Unsicherheiten hin. Studien unterscheiden sich stark bei Partikelgrößen, Polymerarten, Konzentrationen und Nachweismethoden. Für die Risikobewertung ist daher entscheidend, realistische Expositionen von experimentellen Hochdosis-Szenarien zu unterscheiden.""", 2)

add_body("""Die WHO kam in ihrer Bewertung von Mikroplastik im Trinkwasser zu dem Schluss, dass die damalige Datenlage keine einfachen, weitreichenden Gesundheitsbehauptungen zulässt und mehr Forschung benötigt wird. Seitdem ist die Zahl der Studien deutlich gestiegen, doch methodische Standardisierung bleibt ein zentrales Thema. Wissenschaftlich korrekt ist deshalb weder die Aussage „Mikroplastik ist harmlos“ noch die pauschale Aussage „Mikroplastik macht krank“. Besser ist eine abgestufte Formulierung: Menschen sind exponiert; für verschiedene biologische Wirkmechanismen gibt es experimentelle Hinweise; die Größe möglicher Risiken bei typischen Umweltkonzentrationen ist für viele Endpunkte weiterhin Gegenstand der Forschung.""", 0)

add_heading("Was kann man tun?", 2, 3)
add_body("""Minderungsstrategien sollten möglichst an großen und gut beeinflussbaren Quellen ansetzen. Dazu gehören eine bessere Abfallvermeidung und Kreislaufführung von Kunststoffen, geringere Verluste von Kunststoffpellets, abriebärmere Produkte, technische Rückhaltemaßnahmen sowie eine Weiterentwicklung von Waschmaschinen-, Abwasser- und Straßenentwässerungssystemen. Bei Textilien können langlebige Materialien, angepasste Waschgewohnheiten und Filterlösungen die Freisetzung von Fasern beeinflussen. Im Verkehr hängen Reifenabrieb und dessen Eintrag nicht nur vom Reifenmaterial ab, sondern auch von Fahrleistung, Fahrzeugmasse, Fahrstil und Straßenbedingungen. Maßnahmen müssen daher technische und gesellschaftliche Aspekte verbinden.""", 1)

add_body("""Für Verbraucherinnen und Verbraucher ist es sinnvoll, Kunststoffprodukte lange zu nutzen, unnötige Einwegprodukte zu vermeiden, Abfälle korrekt zu entsorgen und beim Kauf auf Langlebigkeit zu achten. Gleichzeitig wäre es falsch, die Verantwortung ausschließlich bei Einzelpersonen abzuladen. Viele relevante Einträge entstehen systemisch und können nur über Produktstandards, Infrastruktur, Industrieprozesse und politische Rahmenbedingungen wirksam reduziert werden. Die europäische Regulierung absichtlich zugesetzter Mikroplastikpartikel ist ein Beispiel dafür, wie konkrete Produktanwendungen schrittweise beschränkt werden können.""", 3)

add_body("""Das Thema Mikroplastik eignet sich deshalb besonders gut, um wissenschaftliches Arbeiten und Dokumentgestaltung gemeinsam zu üben. Eine gute Ausarbeitung trennt Fakten von Bewertungen, kennzeichnet Unsicherheiten, verwendet konsistente Begriffe und gibt Quellen nachvollziehbar an. Ebenso wichtig ist die visuelle Struktur: Überschriften sollten eine erkennbare Hierarchie bilden, Tabellen und Abbildungen benötigen Nummern und Beschriftungen, und das Literaturverzeichnis sollte einem einheitlichen Schema folgen. Genau diese formalen Aspekte sind in der vorliegenden Rohfassung absichtlich uneinheitlich gestaltet.""", 0)

doc.add_page_break()

# Page 6: deliberately inconsistent bibliography
add_heading("Quellen / Literatur / Links", 1, 1)
sources = [
    "1) World Health Organization (2019): Microplastics in drinking-water. https://www.who.int/publications/i/item/9789241516198",
    "NOAA Marine Debris Program: “Microplastics”. https://marinedebris.noaa.gov/what-marine-debris/microplastics (abgerufen 28.09.2026)",
    "UNEP 2020. Microplastics in wastewater: towards solutions – https://www.unep.org/news-and-stories/story/microplastics-wastewater-towards-solutions",
    "European Environment Agency (2021), Microplastics from textiles: towards a circular economy for textiles in Europe. DOI 10.2800/512375",
    "OECD: Global Plastics Outlook, 2022, https://www.oecd.org/en/publications/global-plastics-outlook_de747aef-en.html",
    "Boucher, J.; Friot, D. (2017) Primary microplastics in the oceans: a global evaluation of sources. IUCN. https://iucn.org/resources/publication/primary-microplastics-oceans",
    "ECHA – Mikroplastik. https://echa.europa.eu/de/hot-topics/microplastics",
    "European Parliament. Microplastics: sources, effects and EU solutions (2018). https://www.europarl.europa.eu/topics/en/article/20181116STO19217/mikroplast-kilder-virkninger-og-eu-losninger",
    "Prata JC, da Costa JP, Lopes I, Duarte AC, Rocha-Santos T. Environmental exposure to microplastics: An overview on possible human health effects. Sci Total Environ. 2020;702:134455. doi:10.1016/j.scitotenv.2019.134455",
    "Li Y et al. Potential Health Impact of Microplastics: A Review of Environmental Distribution, Human Exposure, and Toxic Effects. Environ Health. 2023;1(4):249–257. doi:10.1021/envhealth.3c00052",
    "Zuri G, Karanasiou A, Lacorte S. Human biomonitoring of microplastics and health implications: A review. Environ Res. 2023;237:116966. doi:10.1016/j.envres.2023.116966",
    "IUCN (2022): The plastic pollution crisis. https://iucn.org/story/202207/plastic-pollution-crisis",
]
for i, source in enumerate(sources):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3 if i % 3 else 0)
    p.paragraph_format.space_after = Pt([1, 6, 3][i % 3])
    rr = p.add_run(source)
    rr.font.name = ["Arial", "Times New Roman", "Calibri"][i % 3]
    rr.font.size = Pt([9, 10.5, 8.5][i % 3])
    if i in (1, 6):
        rr.underline = True
    if i in (3, 9):
        rr.bold = True

p = doc.add_paragraph("Hinweis: Die Quellenliste ist absichtlich nicht einheitlich formatiert. Genau das gehört zur Übung.")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for rr in p.runs:
    rr.font.name = "Verdana"; rr.font.size = Pt(10); rr.bold = True

doc.save(os.path.join(OUT, "mikroplastik_formatierungsuebung.docx"))
print(os.path.join(OUT, "mikroplastik_formatierungsuebung.docx"))
